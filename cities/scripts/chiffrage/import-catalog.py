#!/usr/bin/env python3
"""Extract private, versioned catalog JSON from the provided Excel sources.

Requires openpyxl. Output must be outside the public Git checkout; it is then
uploaded through the authenticated server import, never committed or deployed
as a static asset. The source files are never copied into the repository.
"""

import argparse
import hashlib
import json
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


def code(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def amount(value):
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace("'", "").replace("’", "").replace(" ", ""))
        except ValueError:
            pass
    return None


def workbook_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def innosolv(source, version):
    workbook = load_workbook(source, read_only=False, data_only=True, keep_vba=True)
    result = []
    dimensions = {"PL_Gemeinde": "population", "PL_Versorger": "meters",
                  "PL_Kirche": "members"}
    defaults = {"1", "27", "30", "100", "116", "125", "129", "135",
                "200", "201", "202", "205", "208", "209", "211", "212", "214"}
    for sheet_name, basis in dimensions.items():
        sheet = workbook[sheet_name]
        anchors = [(col, amount(sheet.cell(1, col).value))
                   for col in range(14, sheet.max_column + 1)]
        anchors = [(col, boundary) for col, boundary in anchors if boundary is not None]
        for row in range(10, sheet.max_row + 1):
            item_code = code(sheet.cell(row, 1).value)
            label = sheet.cell(row, 5).value
            if not item_code or not isinstance(label, str) or not label.strip():
                continue
            tiers = [{"up_to": int(boundary), "value": value}
                     for col, boundary in anchors
                     if (value := amount(sheet.cell(row, col).value)) is not None]
            if not tiers:
                continue
            # Gemeinde is verified against the example calculators. The other
            # product sheets are retained, but unverified rules cannot be used
            # for automatic prices until their source formulas are checked.
            rule = ({"kind": "linear_rounded", "basis": basis, "round_to": 100,
                     "tiers": tiers} if sheet_name == "PL_Gemeinde" else
                    {"kind": "unverified", "basis": basis, "tiers": tiers})
            result.append({"vendor": "innosolv", "product": sheet_name[3:],
                           "item_code": item_code,
                           "label_de": sheet.cell(row, 2).value,
                           "label_fr": label,
                           "description_fr": sheet.cell(row, 7).value,
                           "admin_note_de": sheet.cell(row, 3).value,
                           "admin_note_fr": sheet.cell(row, 6).value,
                           "pricing_rule": rule,
                           "default_selected": sheet_name == "PL_Gemeinde" and
                           (item_code in defaults or item_code == "129VD"),
                           "catalog_version": version,
                           "source": f"{source.name}:{sheet_name}:SW-ID {item_code}",
                           "source_sha256": workbook_hash(source)})
    return result


def abacus(source, sheet_name, flag_col, header_row, version):
    workbook = load_workbook(source, read_only=False, data_only=True)
    sheet = workbook[sheet_name]
    anchors = [(col, amount(sheet.cell(header_row, col).value))
               for col in range(3, 21)]
    anchors = [(col, boundary) for col, boundary in anchors if boundary is not None]
    result = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        item_code = code(sheet.cell(row, 1).value)
        label = sheet.cell(row, 2).value
        if not item_code or not isinstance(label, str):
            continue
        tiers = [{"up_to": int(boundary), "value": value}
                 for col, boundary in anchors
                 if (value := amount(sheet.cell(row, col).value)) is not None]
        if not tiers:
            continue
        result[item_code] = {"vendor": "abacus", "product": "ERP",
                             "item_code": item_code, "label_fr": label,
                             "label_de": None,
                             "rate_class": ("rh_sal_ebanking" if item_code.startswith(
                                 ("10540.", "10515.")) else "standard"),
                             "pricing_rule": {"kind": "lookup", "basis": "population",
                                              "tiers": tiers},
                             "default_selected": str(sheet.cell(row, flag_col).value).lower() == "x",
                             "catalog_version": version,
                             "source": f"{source.name}:{sheet_name}:article {item_code}",
                             "source_sha256": workbook_hash(source)}
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--innosolv", required=True, type=Path)
    parser.add_argument("--avenches", required=True, type=Path)
    parser.add_argument("--fribourg", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--version", required=True)
    parser.add_argument("--abacus-version", required=True,
                        help="Explicit historical Abacus reference version")
    parser.add_argument("--effective-from", type=date.fromisoformat, required=True)
    args = parser.parse_args()
    root = Path(subprocess.check_output(["git", "rev-parse", "--show-toplevel"],
                                        cwd=Path(__file__).parent, text=True).strip()).resolve()
    output = args.out.resolve()
    if output == root or root in output.parents:
        parser.error("Private catalog output must be outside the public Git checkout")
    output.parent.mkdir(parents=True, exist_ok=True)
    items = innosolv(args.innosolv, args.version)
    avenches = abacus(args.avenches, "Abacus", 23, 2, args.abacus_version)
    fribourg = abacus(args.fribourg, "ABA_Client", 25, 5, args.abacus_version)
    for item_code, entry in fribourg.items():
        older = avenches.get(item_code)
        entry["default_selected"] = bool(older and older["default_selected"] and
                                         entry["default_selected"])
        if older and older["pricing_rule"]["tiers"] != entry["pricing_rule"]["tiers"]:
            entry["pricing_rule"]["kind"] = "unverified"
            entry["conflict_sources"] = [older["source"], entry["source"]]
            entry["alternative_tiers"] = older["pricing_rule"]["tiers"]
        items.append(entry)
    output.write_text(json.dumps({"version": args.version,
                                  "effective_from": args.effective_from.isoformat(),
                                  "items": items}, ensure_ascii=False, indent=2) + "\n")
    print(f"Extracted {len(items)} entries to private output; verify unpriced rows before import.")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Prepare private versioned references; never writes to a database or Git.

The output contains confidential business data. Keep every input/output outside
this public checkout. Applying/activating the resulting catalog is separate.
"""
import argparse, hashlib, json, re, subprocess
from pathlib import Path
from openpyxl import load_workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--parameters',required=True,type=Path)
parser.add_argument('--catalog',required=True,type=Path)
parser.add_argument('--workbook',action='append',required=True,type=Path)
parser.add_argument('--default-source',required=True)
parser.add_argument('--output',required=True,type=Path)
args=parser.parse_args()
root=Path(subprocess.check_output(['git','rev-parse','--show-toplevel'],text=True).strip()).resolve()
for path in [args.parameters,args.catalog,*args.workbook,args.output]:
 if path.resolve().is_relative_to(root):raise SystemExit('Private inputs and output must be outside the checkout')
parameters=json.loads(args.parameters.read_text());parameters=parameters.get('content',parameters)
raw=json.loads(args.catalog.read_text());items=[row.get('content',row) for row in (raw.get('items',[]) if isinstance(raw,dict) else raw)]
ids={str(row['item_code']) for row in items if row['vendor']=='innosolv'}
profiles=[]
num=lambda value:isinstance(value,(int,float)) and not isinstance(value,bool)
code=lambda value:str(int(value)) if num(value) and int(value)==value else str(value)
for path in args.workbook:
 w=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True);s=w['Licences&Presta'];values=cached[s.title]
 source=f'{path.name}:Licences&Presta';profile_id=hashlib.sha256(path.name.encode()).hexdigest()[:12]
 label=next((name for name in ['Tolochenaz','Haute-Sorne','Montagny','Fribourg'] if name.lower() in path.name.lower()),path.stem)
 profile={'id':profile_id,'label':f'Référence {label}', 'source':source,'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'default':args.default_source in path.name,'module_days':{},'module_group_labels':{},'services':[],'default_modules':{}}
 days=profile['module_days'];selected=[]
 for row in s.iter_rows(min_row=22):
  r=row[0].row;sw=code(s.cell(r,4).value)
  if sw not in ids:continue
  product='Versorger' if sw in ['501','502','533'] else 'Gemeinde';key=f'innosolv/{product}/{sw}'
  table_days=s.cell(r,12).value
  if num(table_days):days[key]=table_days
  if str(s.cell(r,3).value).lower()=='x':
   if sw not in ['21','22','501','502','533']:selected.append(sw)
   literal=s.cell(r,10).value
   if num(literal):days[key]=literal
 if selected:profile['default_modules']['innosolv']=sorted(set(selected))
 for r in range(6,19):
  label_text=str(s.cell(r,7).value or '')
  service_id=next((key for key,pattern in [('reprise','Reprise de données Prime'),('formation','Formation Prime'),('demarrage','Assistance au démarrage Prime')] if label_text==pattern),None)
  if service_id:
   value=s.cell(r,8).value
   if not num(value):raise ValueError(f'No literal source days at {path.name}:H{r}')
   profile['services'].append({'family':'prestations','item_code':service_id,'label':label_text.replace(' Prime',''),'days':value,'suggested_days':value,'explanation':{'source':f'{source}:H{r}'}})
 gp=next((r for r in range(1,16) if s.cell(r,12).value=='Gestion de projet'),None)
 percent=s.cell(gp,13).value if gp else None
 if not num(percent):raise ValueError(f'No explicit project percentage in {path.name}')
 # Both provided formula forms round GP to one whole day. Refuse other formulas.
 formula=next((str(s.cell(r,c).value) for r in range(6,13) for c in [8,9] if str(s.cell(r,7).value)=='Gestion de projet Prime' and 'MROUND' in str(s.cell(r,c).value)),None)
 if not formula or not ('*$M$7' in formula and ('),1)' in formula or ',$M$2)' in formula)):raise ValueError(f'Unknown GP formula in {path.name}')
 profile['project_management']={'percent':percent*100,'basis':'prime_service_days','round_days':1,'exclude_codes':['deplacements'],'source':f'{source}:M{gp}; {formula}'}
 for r in range(22,36):
  text=str(s.cell(r,5).value or '');value=s.cell(r,8).value
  key=next((key for key,pattern in [('pce-finances','Finances PCE'),('pce-salaires','Salaires PCE'),('installation','Technique (')] if text.startswith(pattern)),None)
  if key and num(value):
   entry={'family':'prestations','item_code':key,'label':text,'days':value,'suggested_days':value,'explanation':{'source':f'{source}:H{r}'}}
   if key.startswith('pce-'):entry.update(requires_products=['pce'],requires_pce_option='finances' if key=='pce-finances' else 'salaires')
   profile['services'].append(entry)
   if key.startswith('pce-'):
    profile.setdefault('pce_defaults',{})['finances' if key=='pce-finances' else 'salaires']=str(s.cell(r,3).value).lower()=='x'
 for row in s.iter_rows(min_row=22,max_row=40):
  r=row[0].row;text=str(s.cell(r,5).value or '');value=s.cell(r,8).value
  if text.startswith('Réserve') and num(value) and str(s.cell(r,3).value).lower()=='x':
   profile['services'].append({'family':'prestations','item_code':'reserve','label':'Réserve de prestations','days':value,'suggested_days':value,'explanation':{'source':f'{source}:H{r}'}})
 # Abacus totals are package implementation days. Allocate once per package,
 # never once per selected option. Blank source cells are not converted to zero.
 if 'ABA_Client' in w.sheetnames:
  a=w['ABA_Client'];av=cached['ABA_Client'];parent=None
  for row in a:
   r=row[0].row;sw=code(a.cell(r,1).value)
   if re.fullmatch(r'105\d+\.1',sw):parent=sw.split('.')[0]+'.'
   text=str(a.cell(r,2).value or '');value=av.cell(r,31).value
   if text.startswith('Total ') and num(value) and parent:
    members=[item for item in items if item['vendor']=='abacus' and str(item['item_code']).startswith(parent)]
    if members:
     profile['services'].append({'family':'prestations','item_code':f'abacus-{parent}','label':text.removeprefix('Total '),'days':value,'suggested_days':value,'requires_products':['abacus'],'requires_module_prefix':f'abacus/ERP/{parent}','explanation':{'source':f'{path.name}:ABA_Client:AE{r}'}})
     for item in members:
      key=f"abacus/ERP/{item['item_code']}";days[key]=0;profile['module_group_labels'][key]=text.removeprefix('Total ')
 if args.default_source in path.name:
  full=values['D16'].value;erp=values['D17'].value;other=values['D18'].value
  if all(num(value) for value in [full,erp,other]):
   parameters['hosting_presets']=[{'label':f'Référence {label}','products':['innosolv','pce'],'lines':[{'family':'prime','category':'hosting','item_code':'hosting-reference','label':f'Hébergement · référence {label}','quantity':1,'investment_pv':full,'annual_pv':erp+other,'explanation':{'source':f'{source}:D16+D17+D18'}}]}]
 profiles.append(profile)
if sum(bool(row['default']) for row in profiles)!=1:raise SystemExit('Exactly one default source is required')
parameters['configuration_profiles']=profiles
args.output.write_text(json.dumps(parameters,ensure_ascii=False,indent=2)+'\n')
print(f'Prepared {len(profiles)} private references. No database writes performed.')
